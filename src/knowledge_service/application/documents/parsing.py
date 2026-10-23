"""将支持的文档格式解析为有序、带来源信息的结构块。"""

import base64
import binascii
import copy
import csv
import hashlib
import io
import json
import mimetypes
import re
import shutil
import subprocess
import tempfile
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field, replace
from datetime import date, datetime
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup
from docx import Document as WordDocument
from pypdf import PdfReader

from knowledge_service.core.config import get_settings
from knowledge_service.infrastructure.ocr_client import (
    LayoutResult,
    OcrServiceError,
    analyze_image,
    analyze_pdf_pages,
    recognize_pdf_pages,
)


class DocumentParseError(ValueError):
    """文档无法安全解析，或需要尚未配置的 OCR 服务。"""


@dataclass(frozen=True)
class ParsedAsset:
    """解析阶段发现的图片；入库阶段负责保存二进制和建立数据库记录。"""

    content: bytes
    media_type: str
    name: str | None = None
    description: str | None = None
    metadata: dict[str, Any] = field(default_factory=dict)

    @property
    def marker(self) -> str:
        """返回解析阶段使用的稳定图片占位符。"""
        value = self.metadata.get("asset_marker")
        return str(value) if value else ""


@dataclass(frozen=True)
class ParsedBlock:
    """保留页码、章节与表格结构的单个解析块。"""

    kind: str
    content: str
    page_number: int | None = None
    section: str | None = None
    metadata: dict[str, Any] = field(default_factory=dict)
    # 图片原始内容由入库阶段写入本地对象目录或 RustFS。
    asset_content: bytes | None = None
    asset_media_type: str | None = None
    # 一个表格块可能包含多个图片单元格，因此支持多个资产。
    assets: tuple[ParsedAsset, ...] = ()


def _asset_marker(content: bytes, name: str, source: str, **location: Any) -> str:
    """为图片生成与解析顺序无关的稳定占位符。"""
    identity = content + name.encode() + source.encode() + json.dumps(
        location, sort_keys=True, ensure_ascii=False
    ).encode()
    return f"[[KB_ASSET:{hashlib.sha256(identity).hexdigest()[:32]}]]"


def _with_asset_marker(text: str, marker: str) -> str:
    """把图片占位符放在原始图片所在文本位置。"""
    return f"{marker}\n{text}" if text.strip() else marker


def parse_document(filename: str, content: bytes) -> list[ParsedBlock]:
    """按文件扩展名选择解析器，并拒绝空文档和无文本扫描件。"""
    extension = Path(filename).suffix.lower()
    parsers = {
        ".txt": _parse_text,
        ".md": _parse_markdown,
        ".html": _parse_html,
        ".htm": _parse_html,
        ".xml": _parse_html,
        ".csv": _parse_csv,
        ".json": _parse_json,
        ".log": _parse_text,
        ".rtf": _parse_rtf,
        ".doc": _parse_doc,
        ".docm": _parse_docx,
        ".docx": _parse_docx,
        ".xls": _parse_xls,
        ".xlsm": _parse_xlsx,
        ".xlsx": _parse_xlsx,
        ".ppt": _parse_ppt,
        ".pptm": _parse_pptx,
        ".pptx": _parse_pptx,
        ".odt": _parse_odt,
        ".ods": _parse_ods,
        ".odp": _parse_odp,
        ".epub": _parse_epub,
        ".pdf": _parse_pdf,
    }
    image_extensions = {".bmp", ".gif", ".jpeg", ".jpg", ".png", ".tif", ".tiff", ".webp"}
    if extension in image_extensions:
        blocks = _parse_image(content, extension)
    else:
        parser = parsers.get(extension)
        if parser is None:
            raise DocumentParseError("UNSUPPORTED_FILE_TYPE: 不支持的文件类型")
        blocks = parser(content)
    blocks = [block for block in blocks if block.content.strip()]
    # 原图说明只参与召回，不作为原文展示；所有格式复用同一个规范化入口。
    normalized: list[ParsedBlock] = []
    for index, block in enumerate(blocks):
        if block.asset_content is not None:
            marker = _asset_marker(block.asset_content, str(index), extension)
            description = re.sub(r"\[\[KB_ASSET:[^\]]+\]\]", "", block.content).strip()
            block = replace(
                block,
                content=marker,
                metadata=block.metadata | {
                    "asset_marker": marker,
                    "asset_description": description,
                    "reading_order": index,
                },
            )
        normalized.append(block)
    blocks = _add_section_path_metadata(normalized)
    if not blocks:
        raise DocumentParseError("EMPTY_DOCUMENT: 文档未提取到可用文本")
    return blocks


def _add_section_path_metadata(blocks: list[ParsedBlock]) -> list[ParsedBlock]:
    """Attach one canonical hierarchy representation while keeping legacy section text."""
    occurrences: dict[tuple[str, str], int] = {}
    previous_path: list[str] = []
    previous_ids: list[str] = []
    normalized: list[ParsedBlock] = []
    for block in blocks:
        raw_path = block.metadata.get("section_path")
        if isinstance(raw_path, (list, tuple)):
            path = [str(part).strip() for part in raw_path if str(part).strip()]
        else:
            path = [
                part.strip()
                for part in str(block.section or "").split(" / ")
                if part.strip()
            ]
        if not path:
            previous_path = []
            previous_ids = []
            normalized.append(
                replace(
                    block,
                    metadata=block.metadata
                    | {
                        "section_path": [],
                        "section_path_ids": [],
                        "heading_schema_version": "heading-path-v1",
                    },
                )
            )
            continue

        common = 0
        while common < min(len(path), len(previous_path)) and path[common] == previous_path[common]:
            common += 1
        path_ids = previous_ids[:common]
        for depth in range(common, len(path)):
            parent_id = path_ids[-1] if path_ids else "root"
            key = (parent_id, path[depth])
            occurrence = occurrences.get(key, 0) + 1
            occurrences[key] = occurrence
            identity = hashlib.sha256(
                f"{parent_id}\0{path[depth]}\0{occurrence}".encode()
            ).hexdigest()[:24]
            path_ids.append(identity)
        previous_path = path
        previous_ids = path_ids
        normalized.append(
            replace(
                block,
                metadata=block.metadata
                | {
                    "section_path": path,
                    "section_path_ids": path_ids,
                    "heading_schema_version": "heading-path-v1",
                },
            )
        )
    return normalized


def _decode_text(content: bytes) -> str:
    """兼容常见中文文本编码，不静默丢弃无法解码的字节。"""
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentParseError("TEXT_ENCODING_ERROR: 文本编码无法识别")


def _clean_text(value: str) -> str:
    """清理空字符与行尾空格，保留段落边界和原始语义。"""
    value = value.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.rstrip() for line in value.split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _docx_heading_level(
    style: str,
    text: str,
    *,
    outline_level: int | None = None,
    allow_numbered_fallback: bool = True,
) -> int | None:
    """识别 Word 标题及常见的无结构编号标题。

    文档作者经常把小节保存成 ``Normal``，因此在没有原生大纲信息时，
    才使用编号推断。编号后必须是短标题且不能以句末标点结束，减少将
    ``①按下按键。`` 之类的操作步骤误判为章节的机会。
    """
    if outline_level is not None and outline_level >= 0:
        return outline_level + 1

    # Word 内置样式在不同语言版本中名称不同（Heading 1、标题 1、标题1）。
    match = re.search(r"(?:heading|标题|überschrift)\s*(\d+)", style or "", re.IGNORECASE)
    if match:
        return int(match.group(1))

    if not allow_numbered_fallback:
        return None

    candidate = _clean_text(text)
    if not candidate or len(candidate) > 100:
        return None
    if re.search(r"[。！？!?；;:.]\s*$", candidate):
        return None

    # 阿拉伯数字、中文序号、括号序号和带圈数字均可作为标题前缀。
    patterns = (
        (r"^\s*(\d{1,3}(?:[.．]\d{1,3}){0,3})\s+(.{1,90})\s*$", "arabic"),
        (r"^\s*([一二三四五六七八九十百千万]+)、\s*(.{1,90})\s*$", "chinese"),
        (r"^\s*[（(]([一二三四五六七八九十百千万]+)[）)]\s*(.{1,90})\s*$", "chinese_sub"),
        (r"^\s*[（(](\d{1,3})[）)]\s*(.{1,90})\s*$", "paren"),
        (r"^\s*([①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳])\s*(.{1,90})\s*$", "circled"),
    )
    for pattern, kind in patterns:
        match = re.match(pattern, candidate)
        if not match:
            continue
        suffix = match.group(2).strip()
        if not suffix or re.search(r"[。！？!?；;:.]", suffix):
            continue
        if kind == "arabic":
            return match.group(1).count(".") + 1
        if kind in {"chinese", "circled"}:
            return 1
        return 2
    return None


def _heading_level(text: str) -> int | None:
    """跨格式的保守编号标题识别入口。"""
    return _docx_heading_level("", text)


def _safe_inferred_heading_level(text: str) -> int | None:
    """Infer a numbered heading only with a structural cue beyond its number."""
    level = _heading_level(text)
    if level is None:
        return None
    deep_numbering = bool(re.match(r"^\s*\d{1,3}(?:[.．]\d{1,3}){2,}(?:\s|$)", text))
    heading_cues = re.compile(
        r"方法|步骤|注意事项|清洁|维护|故障|原理|简介|概述|适用范围|产品介绍|"
        r"功能|配件|规格|参数|设置|准备|附录|参考文献|使用说明|操作说明|"
        r"\b(method|steps?|procedure|overview|introduction|purpose|scope|"
        r"maintenance|precautions|specifications|appendix)\b",
        re.IGNORECASE,
    )
    return level if deep_numbering or heading_cues.search(text) else None


def _section_path(headings: list[tuple[int, str]]) -> str | None:
    return " / ".join(title for _, title in headings) or None


def _parse_text(content: bytes) -> list[ParsedBlock]:
    text = _clean_text(_decode_text(content))
    headings: list[tuple[int, str]] = []
    blocks: list[ParsedBlock] = []
    for part in re.split(r"\n\s*\n", text):
        level = _safe_inferred_heading_level(part)
        if level is not None:
            headings = [entry for entry in headings if entry[0] < level]
            headings.append((level, part))
        blocks.append(ParsedBlock(kind="text", content=part, section=_section_path(headings)))
    return blocks


def _markdown_table_cells(line: str) -> list[str] | None:
    """解析 GFM 表格行；不是表格行时返回 None。"""
    stripped = line.strip()
    if not (stripped.startswith("|") and stripped.endswith("|")):
        return None
    cells = re.split(r"(?<!\\)\|", stripped[1:-1])
    return [cell.strip().replace("\\|", "|") for cell in cells]


def _is_markdown_table_delimiter(cells: list[str]) -> bool:
    """分隔行单元格形如 ---、:---、---:、:---:。"""
    return all(re.fullmatch(r":?-{1,}:?", cell) for cell in cells)


def _parse_markdown(content: bytes) -> list[ParsedBlock]:
    """识别 ATX/Setext 标题与 GFM 表格，忽略代码围栏中的标题语法。"""
    text = _clean_text(_decode_text(content))
    headings: list[tuple[int, str]] = []
    blocks: list[ParsedBlock] = []
    body_lines: list[str] = []
    fence: tuple[str, int] | None = None
    lines = text.splitlines()
    index = 0

    def flush() -> None:
        body = _clean_text("\n".join(body_lines))
        body_lines.clear()
        if not body:
            return
        blocks.append(
            ParsedBlock(
                kind="text",
                content=body,
                section=" / ".join(x[1] for x in headings) or None,
            )
        )

    while index < len(lines):
        line = lines[index]
        fence_match = re.match(r"^\s{0,3}(`{3,}|~{3,})", line)
        if fence:
            body_lines.append(line)
            if (
                fence_match
                and fence_match.group(1)[0] == fence[0]
                and len(fence_match.group(1)) >= fence[1]
            ):
                fence = None
            index += 1
            continue
        if fence_match:
            fence = (fence_match.group(1)[0], len(fence_match.group(1)))
            body_lines.append(line)
            index += 1
            continue
        cells = _markdown_table_cells(line)
        if cells is not None:
            # 表格行 + 分隔行才构成 GFM 表格；仅一行带竖线的文本按正文处理。
            next_cells = (
                _markdown_table_cells(lines[index + 1]) if index + 1 < len(lines) else None
            )
            if next_cells is not None and _is_markdown_table_delimiter(next_cells):
                flush()
                rows = [cells]
                index += 2
                while index < len(lines):
                    row_cells = _markdown_table_cells(lines[index])
                    if row_cells is None or not lines[index].strip():
                        break
                    rows.append(row_cells)
                    index += 1
                header = " | ".join(rows[0])
                blocks.append(
                    ParsedBlock(
                        kind="table",
                        content="\n".join(
                            [header, *(" | ".join(row) for row in rows[1:])]
                        ),
                        section=" / ".join(x[1] for x in headings) or None,
                        metadata={"table_header": header},
                    )
                )
                continue
        atx = re.match(r"^\s{0,3}(#{1,6})\s+(.+?)\s*#*\s*$", line)
        if atx:
            flush()
            level, title = len(atx.group(1)), atx.group(2).strip()
            headings[:] = [item for item in headings if item[0] < level]
            headings.append((level, title))
            index += 1
            continue
        setext = re.match(r"^\s{0,3}(=+|-+)\s*$", line)
        if setext and body_lines:
            title_lines = [part.strip() for part in body_lines if part.strip()]
            body_lines.clear()
            title = _clean_text(" ".join(title_lines))
            if title:
                level = 1 if setext.group(1).startswith("=") else 2
                headings[:] = [item for item in headings if item[0] < level]
                headings.append((level, title))
            index += 1
            continue
        if not line.strip():
            flush()
        else:
            body_lines.append(line)
        index += 1
    flush()
    return blocks


def _parse_html(content: bytes) -> list[ParsedBlock]:
    """去除页面脚本等非正文内容，同时保留标题层级和 HTML 表格。"""
    soup = BeautifulSoup(content, "html.parser")
    for node in soup(["script", "style", "noscript", "svg"]):
        node.decompose()
    root = soup.body or soup
    blocks: list[ParsedBlock] = []
    headings: list[tuple[int, str]] = []
    for block_index, node in enumerate(
        root.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "table"])
    ):
        if node.name != "table" and node.find_parent("table"):
            continue
        if node.name == "p" and node.find_parent("li"):
            continue
        if node.name.startswith("h"):
            level = int(node.name[1])
            title = _clean_text(node.get_text(" ", strip=True))
            headings = [item for item in headings if item[0] < level]
            headings.append((level, title))
        elif node.name == "table":
            rows: list[list[str]] = []
            table_assets: list[ParsedAsset] = []
            for row_index, row in enumerate(node.find_all("tr")):
                values = []
                for column_index, cell in enumerate(row.find_all(["th", "td"])):
                    cell_assets: list[ParsedAsset] = []
                    values.append(
                        _html_inline_content(
                            cell,
                            cell_assets,
                            location={
                                "format": "html_table",
                                "block_index": block_index,
                                "row": row_index,
                                "column": column_index,
                            },
                        )
                    )
                    table_assets.extend(cell_assets)
                rows.append(values)
            rows = [row for row in rows if any(row)]
            if rows:
                header = rows[0]
                rendered = [" | ".join(header)]
                rendered.extend(" | ".join(row) for row in rows[1:])
                blocks.append(
                    ParsedBlock(
                        kind="table",
                        content="\n".join(rendered),
                        section=" / ".join(x[1] for x in headings) or None,
                        metadata={"table_header": " | ".join(header)},
                        assets=tuple(table_assets),
                    )
                )
        else:
            inline_assets: list[ParsedAsset] = []
            text = _html_inline_content(
                node,
                inline_assets,
                location={"format": "html", "block_index": block_index},
            )
            if text:
                blocks.append(
                    ParsedBlock(
                        kind="text",
                        content=text,
                        section=" / ".join(x[1] for x in headings) or None,
                        assets=tuple(inline_assets),
                    )
                )
    if not blocks:
        return [ParsedBlock(kind="text", content=_clean_text(root.get_text("\n", strip=True)))]
    return blocks


def _html_inline_content(
    node: Any,
    assets: list[ParsedAsset] | None = None,
    *,
    location: dict[str, Any] | None = None,
) -> str:
    """保留 HTML 段落或表格单元格中的图片 Markdown 位置与换行。"""
    parts: list[str] = []
    for item in node.descendants:
        if getattr(item, "name", None) == "br":
            parts.append("\n")
        elif getattr(item, "name", None) == "img":
            source = str(item.get("src") or item.get("data-src") or "").strip()
            if source:
                alt = str(item.get("alt") or "").replace("]", "\\]")
                match = re.fullmatch(r"data:(image/[A-Za-z0-9.+-]+);base64,(.+)", source)
                if match and assets is not None:
                    try:
                        image_bytes = base64.b64decode(match.group(2), validate=True)
                    except (ValueError, binascii.Error):
                        parts.append(f"![{alt}]({source})")
                    else:
                        marker = _asset_marker(
                            image_bytes, alt or "html-image", "html", **(location or {})
                        )
                        assets.append(
                            ParsedAsset(
                                image_bytes,
                                match.group(1),
                                alt or None,
                                alt or None,
                                (location or {}) | {"asset_marker": marker},
                            )
                        )
                        parts.append(marker)
                else:
                    parts.append(f"![{alt}]({source})")
        elif isinstance(item, str):
            parts.append(str(item))
    return _clean_text("".join(parts))


def _parse_csv(content: bytes) -> list[ParsedBlock]:
    """将 CSV 记录转成可读行，分块时会为每段重复列名。"""
    text = _decode_text(content)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    rows = [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text), dialect)]
    rows = [row for row in rows if any(row)]
    if len(rows) < 2:
        return [ParsedBlock(kind="text", content="\n".join(" | ".join(row) for row in rows))]
    header = " | ".join(rows[0])
    return [
        ParsedBlock(
            kind="table",
            content="\n".join([header, *(" | ".join(row) for row in rows[1:])]),
            metadata={"table_header": header},
        )
    ]


def _parse_json(content: bytes) -> list[ParsedBlock]:
    """格式化 JSON，避免把原始结构压成难以检索的一行。"""
    try:
        value = json.loads(_decode_text(content))
    except (json.JSONDecodeError, DocumentParseError) as exc:
        raise DocumentParseError("JSON_PARSE_ERROR: JSON 文件无法解析") from exc
    return [
        ParsedBlock(
            kind="text",
            content=json.dumps(value, ensure_ascii=False, indent=2),
            metadata={"format": "json"},
        )
    ]


def _parse_rtf(content: bytes) -> list[ParsedBlock]:
    """提取 RTF 正文，保留段落边界供后续分块。"""
    try:
        from striprtf.striprtf import rtf_to_text

        text = rtf_to_text(_decode_text(content))
    except Exception as exc:
        raise DocumentParseError("RTF_PARSE_ERROR: RTF 文件无法解析") from exc
    cleaned = _clean_text(text)
    return [ParsedBlock(kind="text", content=part) for part in re.split(r"\n\s*\n", cleaned)]


def _format_cell(value: Any) -> str:
    """将办公表格单元格转换为稳定、可检索的文本。"""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _clean_text(str(value))


def _normalize_table_row(values: Any) -> list[str]:
    """清理表格行尾空单元格，避免生成大量无意义分隔符。"""
    row = [_format_cell(value) for value in values]
    while row and not row[-1]:
        row.pop()
    return row


def _table_block(
    rows: list[list[str]], *, section: str, metadata: dict[str, Any]
) -> ParsedBlock | None:
    """把一个工作表或演示文稿表格转换为统一的表格块。"""
    rows = [row for row in rows if any(row)]
    if not rows:
        return None
    header = " | ".join(rows[0])
    return ParsedBlock(
        kind="table",
        content="\n".join(" | ".join(row) for row in rows),
        section=section or None,
        metadata={**metadata, "table_header": header},
    )


def _parse_xlsx(content: bytes) -> list[ParsedBlock]:
    """按工作表读取 XLSX，保留工作表名称和行列结构。"""
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    except Exception as exc:
        raise DocumentParseError("XLSX_PARSE_ERROR: Excel 文件无法解析") from exc
    blocks: list[ParsedBlock] = []
    description_cache: dict[str, str] = {}
    try:
        image_jobs: list[tuple[bytes, str, str]] = []
        for worksheet in workbook.worksheets:
            for image in getattr(worksheet, "_images", []):
                image_format = str(getattr(image, "format", None) or "png").lower()
                image_name = f"{worksheet.title}.{image_format}"
                image_jobs.append(
                    (
                        _xlsx_image_bytes(image),
                        image_name,
                        mimetypes.guess_type(image_name)[0] or "image/png",
                    )
                )
        _prewarm_office_image_descriptions(image_jobs, description_cache, "xlsx")
        for worksheet in workbook.worksheets:
            # 在图片锚点处理完之前保留空行，防止原始行号发生偏移。
            rows = [
                _normalize_table_row(values)
                for values in worksheet.iter_rows(values_only=True)
            ]
            table_assets: list[ParsedAsset] = []
            # Excel 图片通常是浮动对象，通过 anchor 记录左上角所在的行列。
            for image in getattr(worksheet, "_images", []):
                anchor = getattr(image, "anchor", None)
                origin = getattr(anchor, "_from", None)
                if origin is None:
                    continue
                row_index = int(origin.row)
                column_index = int(origin.col)
                while len(rows) <= row_index:
                    rows.append([])
                while len(rows[row_index]) <= column_index:
                    rows[row_index].append("")
                image_format = str(getattr(image, "format", None) or "png").lower()
                image_name = (
                    f"{worksheet.title}-r{row_index + 1}-c{column_index + 1}.{image_format}"
                )
                media_type = mimetypes.guess_type(image_name)[0] or "image/png"
                asset = _describe_office_image(
                    _xlsx_image_bytes(image),
                    image_name,
                    media_type,
                    {
                        "source": "xlsx_cell_image",
                        "sheet": worksheet.title,
                        "row": row_index + 1,
                        "column": column_index + 1,
                    },
                    description_cache,
                )
                table_assets.append(asset)
                rows[row_index][column_index] = _with_asset_marker(
                    rows[row_index][column_index], asset.marker
                )
            block = _table_block(
                rows,
                section=worksheet.title,
                metadata={"sheet": worksheet.title},
            )
            if block:
                blocks.append(
                    ParsedBlock(
                        kind=block.kind,
                        content=block.content,
                        section=block.section,
                        metadata=block.metadata,
                        assets=tuple(table_assets),
                    )
                )
    finally:
        workbook.close()
    return blocks


def _xlsx_image_bytes(image: Any) -> bytes:
    """缓存 Excel 图片二进制，避免预识别关闭流后无法再次提取。"""
    if not hasattr(image, "_kb_content"):
        image._kb_content = bytes(image._data())
    return image._kb_content


def _parse_xls(content: bytes) -> list[ParsedBlock]:
    """按工作表读取旧版 XLS，使用 xlrd 处理 OLE 二进制格式。"""
    try:
        import xlrd

        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
    except Exception as exc:
        raise DocumentParseError("XLS_PARSE_ERROR: 旧版 Excel 文件无法解析") from exc
    blocks: list[ParsedBlock] = []
    for sheet_name in workbook.sheet_names():
        worksheet = workbook.sheet_by_name(sheet_name)
        rows = [
            row
            for row in (
                _normalize_table_row(worksheet.row_values(row_index))
                for row_index in range(worksheet.nrows)
            )
            if row
        ]
        block = _table_block(rows, section=sheet_name, metadata={"sheet": sheet_name})
        if block:
            blocks.append(block)
    workbook.release_resources()
    return blocks


def _parse_pptx(content: bytes) -> list[ParsedBlock]:
    """按幻灯片的版面顺序读取文字、表格和内嵌图片。"""
    try:
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(content))
    except Exception as exc:
        raise DocumentParseError("PPTX_PARSE_ERROR: PowerPoint 文件无法解析") from exc
    blocks: list[ParsedBlock] = []
    image_jobs: list[tuple[bytes, str, str]] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            if getattr(shape, "shape_type", None) == 13 and getattr(shape, "image", None):
                image = shape.image
                image_name = str(
                    getattr(image, "filename", None) or f"image-{slide_number}.png"
                )
                image_jobs.append(
                    (
                        bytes(image.blob),
                        image_name,
                        str(getattr(image, "content_type", None) or "image/png"),
                    )
                )
    description_cache: dict[str, str] = {}
    _prewarm_office_image_descriptions(image_jobs, description_cache, "pptx")
    for slide_number, slide in enumerate(presentation.slides, start=1):
        section = f"第 {slide_number} 页"
        # python-pptx 的 shapes 顺序不一定等于阅读顺序，按顶部再按左侧位置排序。
        shapes = sorted(slide.shapes, key=lambda item: (item.top, item.left))
        for reading_order, shape in enumerate(shapes):
            location = {
                "format": "pptx",
                "slide_number": slide_number,
                "reading_order": reading_order,
                "bbox_emu": [shape.left, shape.top, shape.width, shape.height],
            }
            if getattr(shape, "has_table", False):
                rows = [
                    _normalize_table_row(cell.text for cell in row.cells)
                    for row in shape.table.rows
                ]
                block = _table_block(
                    rows,
                    section=section,
                    metadata={"slide_number": slide_number},
                )
                if block:
                    blocks.append(
                        ParsedBlock(
                            kind=block.kind,
                            content=block.content,
                            page_number=slide_number,
                            section=block.section,
                            metadata=block.metadata,
                        )
                    )
            elif getattr(shape, "has_text_frame", False):
                text = _clean_text(shape.text)
                if text:
                    blocks.append(
                        ParsedBlock(
                            kind="text",
                            content=text,
                            page_number=slide_number,
                            section=section,
                            metadata={"slide_number": slide_number},
                        )
                    )
            elif getattr(shape, "shape_type", None) == 13 and getattr(
                shape, "image", None
            ) is not None:
                blocks.extend(
                    _parse_office_image(
                        bytes(shape.image.blob),
                        str(getattr(shape.image, "filename", None) or f"image-{slide_number}.png"),
                        str(getattr(shape.image, "content_type", None) or "image/png"),
                        section,
                        slide_number,
                        "pptx",
                        description_cache=description_cache,
                        location=location,
                    )
                )
    return blocks


def _parse_office_image(
    content: bytes,
    name: str,
    media_type: str,
    section: str,
    page_number: int | None,
    source: str,
    description_cache: dict[str, str] | None = None,
    location: dict[str, Any] | None = None,
) -> list[ParsedBlock]:
    """统一保存 Office 原图；视觉说明只作为独立的召回信息。"""
    asset = _describe_office_image(
        content, name, media_type,
        {
            "source": f"{source}_embedded_image",
            "page": page_number,
            "section": section,
            **(location or {}),
        },
        description_cache,
    )
    return [
        ParsedBlock(
            kind="text",
            content=asset.description or asset.marker,
            page_number=page_number,
            section=section or None,
            metadata=asset.metadata | {"asset_name": name},
            asset_content=content,
            asset_media_type=media_type,
        )
    ]


def _office_ocr_settings(settings: Any) -> Any:
    """复制配置并单独设置 Office 图片 OCR 超时，不影响 PDF OCR 超时。"""
    result = copy.copy(settings)
    result.ocr_timeout_seconds = float(
        getattr(settings, "ocr_embedded_timeout_seconds", 60)
    )
    return result


def _describe_office_image(
    content: bytes,
    name: str,
    media_type: str,
    metadata: dict[str, Any],
    description_cache: dict[str, str] | None = None,
) -> ParsedAsset:
    """识别 Office 图片并返回带原图和来源位置的资产对象。"""
    settings = get_settings()
    digest = hashlib.sha256(content).hexdigest()
    marker = _asset_marker(content, name, str(metadata.get("source", "office")), **{
        key: value for key, value in metadata.items()
        if key in {"row", "column", "page", "slide", "paragraph", "inline_order"}
    })
    description = description_cache.get(digest) if description_cache is not None else None
    if description is not None:
        return ParsedAsset(
            content, media_type, name, description, metadata | {"asset_marker": marker}
        )
    if not getattr(settings, "ocr_embedded_images_enabled", True):
        description = f"图片：{name}"
    elif not _should_ocr_embedded_image(content, media_type, settings):
        description = f"图片：{name}（小图标或低分辨率图片，未进行文字识别）"
    else:
        try:
            layout = analyze_image(
                content,
                settings=_office_ocr_settings(settings),
                filename=name,
            )
        except OcrServiceError:
            location = f"所在位置：{metadata.get('section')}。" if metadata.get("section") else ""
            description = (
                f"图片说明暂不可用。{location}图片文件名：{name}。原图已保存，"
                "图中内容尚未识别，请勿据此推断具体文字或操作。"
            )
        else:
            description = layout.markdown or "\n".join(item.text for item in layout.blocks)
            description = _clean_text(description) or f"图片：{name}"
    if description_cache is not None:
        description_cache[digest] = description
    return ParsedAsset(
        content=content,
        media_type=media_type,
        name=name,
        description=description,
        metadata=metadata | {"asset_marker": marker},
    )


def _should_ocr_embedded_image(content: bytes, media_type: str, settings: Any) -> bool:
    """按图片类型、大小和分辨率过滤不适合视觉识别的 Office 图片。"""
    if not media_type.startswith("image/") or media_type in {"image/wmf", "image/emf"}:
        return False
    if len(content) < int(getattr(settings, "ocr_embedded_image_min_bytes", 2048)):
        return False
    try:
        from PIL import Image

        with Image.open(io.BytesIO(content)) as image:
            return image.width >= int(
                getattr(settings, "ocr_embedded_image_min_width", 128)
            ) and image.height >= int(
                getattr(settings, "ocr_embedded_image_min_height", 128)
            )
    except Exception:
        # 无法确认尺寸时不把二进制 Office 图形送给图片接口。
        return False


def _docx_cell_images(document: Any, cell: Any) -> list[tuple[bytes, str, str]]:
    """通过 DOCX 关系表找到单元格内嵌图片的二进制内容。"""
    from docx.oxml.ns import qn

    images: list[tuple[bytes, str, str]] = []
    for element in cell._tc.iter():
        if element.tag != qn("a:blip"):
            continue
        relationship_id = element.get(qn("r:embed"))
        related = document.part.related_parts.get(relationship_id)
        if related is None or not hasattr(related, "blob"):
            continue
        part_name = Path(str(related.partname)).name
        media_type = str(getattr(related, "content_type", None) or "image/png")
        images.append((bytes(related.blob), part_name, media_type))
    return images


def _docx_paragraph_events(
    document: Any, paragraph: Any
) -> list[tuple[str, str | tuple[bytes, str, str]]]:
    """按 Word 段落 XML 顺序返回文字和内嵌图片事件。"""
    from docx.oxml.ns import qn

    events: list[tuple[str, str | tuple[bytes, str, str]]] = []
    for run in paragraph.runs:
        for element in run._r.iter():
            if element.tag == qn("w:t") and element.text:
                events.append(("text", element.text))
            elif element.tag == qn("a:blip"):
                relationship_id = element.get(qn("r:embed"))
                related = document.part.related_parts.get(relationship_id)
                if related is None or not hasattr(related, "blob"):
                    continue
                name = Path(str(related.partname)).name
                media_type = str(getattr(related, "content_type", None) or "image/png")
                events.append(("image", (bytes(related.blob), name, media_type)))
    return events


def _prewarm_office_image_descriptions(
    images: list[tuple[bytes, str, str]],
    description_cache: dict[str, str],
    source: str,
) -> None:
    """按小批量并发识别 Office 图片，并按摘要去重结果。"""
    settings = get_settings()
    unique_images: dict[str, tuple[bytes, str, str]] = {}
    for content, name, media_type in images:
        if media_type.startswith("image/"):
            digest = hashlib.sha256(content).hexdigest()
            unique_images.setdefault(digest, (content, name, media_type))
    if not unique_images:
        return
    workers = max(1, min(int(getattr(settings, "ocr_embedded_concurrency", 3)), 8))
    image_items = list(unique_images.items())
    with ThreadPoolExecutor(max_workers=workers, thread_name_prefix="office-ocr") as pool:
        # 分批提交，既限制模型侧并发，也避免一次向线程池塞入大量待处理图片。
        for start in range(0, len(image_items), workers):
            batch = image_items[start : start + workers]
            futures = {
                pool.submit(
                    _describe_office_image,
                    content,
                    name,
                    media_type,
                    {"source": f"{source}_embedded_image", "asset_name": name},
                    {},
                ): digest
                for digest, (content, name, media_type) in batch
            }
            for future in as_completed(futures):
                digest = futures[future]
                try:
                    description_cache[digest] = future.result().description or "图片"
                except Exception:
                    # 单张图片异常不阻断整份文档，其原图仍由后续解析流程保存。
                    description_cache.pop(digest, None)


def _prewarm_docx_image_descriptions(
    document: Any, description_cache: dict[str, str]
) -> None:
    """收集 DOCX 关系图片后交给统一的小批量识别器。"""
    images = [
        (
            bytes(related.blob),
            Path(str(getattr(related, "partname", "image.png"))).name,
            str(getattr(related, "content_type", None) or "image/png"),
        )
        for related in document.part.related_parts.values()
        if hasattr(related, "blob")
    ]
    _prewarm_office_image_descriptions(images, description_cache, "docx")


def _parse_image(content: bytes, extension: str) -> list[ParsedBlock]:
    """调用千帆视觉模型识别常见图片文件中的文字。"""
    try:
        layout = analyze_image(
            content,
            settings=get_settings(),
            filename=f"document{extension}",
        )
    except OcrServiceError as exc:
        raise DocumentParseError(str(exc)) from exc
    blocks = _layout_blocks(
        layout,
        metadata={"format": "image"},
        image_content=content,
        image_media_type=mimetypes.guess_type(f"image{extension}")[0],
    )
    if blocks and not any(block.asset_content is not None for block in blocks):
        marker = _asset_marker(content, f"document{extension}", "image")
        blocks[0] = replace(
            blocks[0],
            content=_with_asset_marker(blocks[0].content, marker),
            metadata=blocks[0].metadata | {"asset_marker": marker},
            asset_content=content,
            asset_media_type=mimetypes.guess_type(f"image{extension}")[0],
        )
    return blocks


def _layout_blocks(
    layout: LayoutResult,
    *,
    page_number: int | None = None,
    section: str | None = None,
    metadata: dict[str, Any] | None = None,
    image_content: bytes | None = None,
    image_media_type: str | None = None,
    page_image: bytes | None = None,
    headings_state: list[tuple[int, str]] | None = None,
) -> list[ParsedBlock]:
    """把视觉模型的文字、表格和图片说明映射为带来源的解析块。"""
    common = metadata or {}
    blocks: list[ParsedBlock] = []
    headings = headings_state if headings_state is not None else []
    for item_index, item in enumerate(layout.blocks):
        item_metadata = common | {"layout_kind": item.kind}
        asset_content = image_content if item.kind == "image" else None
        if item.kind == "image" and page_image and item.bbox:
            try:
                import fitz

                image_doc = fitz.open(stream=page_image, filetype="pdf")
                page = image_doc[(page_number or 1) - 1]
                x0, y0, x1, y1 = item.bbox
                rect = fitz.Rect(
                    page.rect.width * x0 / 1000,
                    page.rect.height * y0 / 1000,
                    page.rect.width * x1 / 1000,
                    page.rect.height * y1 / 1000,
                ) & page.rect
                if not rect.is_empty:
                    asset_content = page.get_pixmap(clip=rect, alpha=False).tobytes("png")
                image_doc.close()
                image_media_type = "image/png"
            except Exception:
                asset_content = None
        if item.bbox is not None:
            item_metadata["bbox"] = list(item.bbox)
        item_text = item.text
        explicit_level = item.metadata.get("heading_level", item.metadata.get("level"))
        try:
            explicit_level = int(explicit_level) if explicit_level is not None else None
        except (TypeError, ValueError):
            explicit_level = None
        inferred_level = explicit_level or (
            _heading_level(item_text)
            if item.kind == "heading"
            else _safe_inferred_heading_level(item_text)
        )
        if item.kind == "heading" and inferred_level is None:
            inferred_level = 1
        if inferred_level is not None and item.kind in {"text", "heading"}:
            headings = [entry for entry in headings if entry[0] < inferred_level]
            headings.append((inferred_level, item_text))
        item_section = _section_path(headings) or section
        if asset_content is not None:
            marker = _asset_marker(
                asset_content,
                f"page-{page_number or 0}.png",
                str(common.get("format", "layout")),
                page=page_number or 0,
                index=item_index,
            )
            item_metadata["asset_marker"] = marker
            item_text = _with_asset_marker(item_text, marker)
        blocks.append(
            ParsedBlock(
                kind="table" if item.kind == "table" else "text",
                content=item_text,
                page_number=page_number,
                section=item_section,
                metadata=item_metadata,
                asset_content=asset_content,
                asset_media_type=image_media_type if item.kind == "image" else None,
            )
        )
    if not blocks and layout.markdown:
        blocks.append(
            ParsedBlock(
                kind="text",
                content=layout.markdown,
                page_number=page_number,
                metadata=common,
            )
        )
    return blocks


def _find_office_converter() -> str | None:
    """查找配置指定或系统 PATH 中的 LibreOffice 转换程序。"""
    settings = get_settings()
    configured = str(settings.office_converter_path).strip()
    if configured:
        return configured
    return shutil.which("soffice") or shutil.which("libreoffice")


def _convert_legacy_office(content: bytes, source_extension: str, target_extension: str) -> bytes:
    """用 LibreOffice 将旧版 Office 二进制文件转换为 OpenXML 后再解析。"""
    converter = _find_office_converter()
    if not converter:
        raise DocumentParseError(
            "OFFICE_CONVERTER_NOT_CONFIGURED: .doc/.ppt 等旧版 Office 文件需要安装 LibreOffice，"
            "并配置 KNOWLEDGE_SERVICE_OFFICE_CONVERTER_PATH"
        )
    settings = get_settings()
    try:
        with tempfile.TemporaryDirectory(prefix="knowledge-office-") as directory:
            workdir = Path(directory)
            source = workdir / f"source{source_extension}"
            source.write_bytes(content)
            profile = workdir / "profile"
            profile.mkdir()
            command = [
                converter,
                "--headless",
                f"-env:UserInstallation={profile.as_uri()}",
                "--convert-to",
                target_extension.removeprefix("."),
                "--outdir",
                str(workdir),
                str(source),
            ]
            result = subprocess.run(
                command,
                capture_output=True,
                text=True,
                timeout=float(settings.office_converter_timeout_seconds),
                check=False,
            )
            output = workdir / f"source{target_extension}"
            if result.returncode != 0 or not output.is_file():
                raise DocumentParseError(
                    "OFFICE_CONVERSION_ERROR: Office 文件转换失败，请检查 LibreOffice 配置"
                )
            return output.read_bytes()
    except subprocess.TimeoutExpired as exc:
        raise DocumentParseError("OFFICE_CONVERSION_TIMEOUT: Office 文件转换超时") from exc
    except OSError as exc:
        raise DocumentParseError("OFFICE_CONVERTER_ERROR: Office 转换程序无法启动") from exc


def _parse_doc(content: bytes) -> list[ParsedBlock]:
    """转换旧版 DOC 后复用 DOCX 解析器。"""
    return _parse_docx(_convert_legacy_office(content, ".doc", ".docx"))


def _parse_ppt(content: bytes) -> list[ParsedBlock]:
    """转换旧版 PPT 后复用 PPTX 解析器。"""
    return _parse_pptx(_convert_legacy_office(content, ".ppt", ".pptx"))


def _parse_odt(content: bytes) -> list[ParsedBlock]:
    """将 OpenDocument 文本文档转换后复用 DOCX 解析器。"""
    return _parse_docx(_convert_legacy_office(content, ".odt", ".docx"))


def _parse_ods(content: bytes) -> list[ParsedBlock]:
    """将 OpenDocument 表格转换后复用 XLSX 解析器。"""
    return _parse_xlsx(_convert_legacy_office(content, ".ods", ".xlsx"))


def _parse_odp(content: bytes) -> list[ParsedBlock]:
    """将 OpenDocument 演示文稿转换后复用 PPTX 解析器。"""
    return _parse_pptx(_convert_legacy_office(content, ".odp", ".pptx"))


def _epub_spine_hrefs(archive: zipfile.ZipFile) -> list[str]:
    """从 container.xml 指向的 OPF 中解析 spine 阅读顺序。"""
    container_name: str | None = None
    try:
        container = BeautifulSoup(archive.read("META-INF/container.xml"), "xml")
    except (KeyError, zipfile.BadZipFile):
        return []
    rootfile = container.find("rootfile")
    if rootfile is not None:
        container_name = rootfile.get("full-path")
    if not container_name:
        return []
    try:
        opf = BeautifulSoup(archive.read(container_name), "xml")
    except (KeyError, zipfile.BadZipFile):
        return []
    manifest = {
        item.get("id"): item.get("href")
        for item in opf.find_all("item")
        if item.get("id") and item.get("href")
    }
    base_dir = Path(container_name).parent
    hrefs: list[str] = []
    for itemref in opf.find_all("itemref"):
        href = manifest.get(itemref.get("idref"))
        if not href:
            continue
        href = href.split("#", 1)[0]
        hrefs.append(str((base_dir / href).as_posix()))
    return hrefs


def _epub_chapter_title(chapter: bytes) -> str | None:
    """取章节内第一个 h1/h2/title 作为章节标题。"""
    soup = BeautifulSoup(chapter, "html.parser")
    for tag in ("h1", "h2", "title"):
        node = soup.find(tag)
        if node is not None:
            title = _clean_text(node.get_text(" ", strip=True))
            if title:
                return title
    return None


def _parse_epub(content: bytes) -> list[ParsedBlock]:
    """按 spine 顺序解析 EPUB 各章节 XHTML，章节标题并入章节路径。"""
    try:
        archive = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile as exc:
        raise DocumentParseError("EPUB_PARSE_ERROR: EPUB 文件无法解析") from exc
    try:
        spine_hrefs = _epub_spine_hrefs(archive)
        if not spine_hrefs:
            # 无有效 OPF 时退回按文件名收集章节。
            spine_hrefs = sorted(
                name
                for name in archive.namelist()
                if name.lower().endswith((".xhtml", ".html", ".htm"))
            )
        blocks: list[ParsedBlock] = []
        chapter_number = 0
        for href in spine_hrefs:
            try:
                chapter = archive.read(href)
            except KeyError:
                continue
            if not chapter.strip():
                continue
            chapter_number += 1
            chapter_blocks = _parse_html(chapter)
            chapter_title = _epub_chapter_title(chapter) or f"第 {chapter_number} 章"
            for block in chapter_blocks:
                block_section = block.section or ""
                if block_section == chapter_title or block_section.startswith(
                    chapter_title + " / "
                ):
                    blocks.append(block)
                    continue
                section = (
                    f"{chapter_title} / {block_section}" if block_section else chapter_title
                )
                blocks.append(replace(block, section=section))
        return blocks
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError("EPUB_PARSE_ERROR: EPUB 文件无法解析") from exc
    finally:
        archive.close()


def _parse_docx(content: bytes) -> list[ParsedBlock]:
    """按文档原始顺序读取 Word 段落和表格。"""
    try:
        document = WordDocument(io.BytesIO(content))
    except Exception as exc:
        raise DocumentParseError("DOCX_PARSE_ERROR: Word 文档无法解析") from exc
    blocks: list[ParsedBlock] = []
    headings: list[tuple[int, str]] = []
    table_image_names: set[str] = set()
    referenced_image_names: set[str] = set()
    description_cache: dict[str, str] = {}
    if getattr(get_settings(), "ocr_embedded_images_enabled", True):
        # 先小批量并发识别，再按 Word 原始结构顺序组装正文，避免串行等待。
        _prewarm_docx_image_descriptions(document, description_cache)
    for item in document.iter_inner_content():
        if hasattr(item, "rows"):
            rows: list[list[str]] = []
            table_assets: list[ParsedAsset] = []
            for row_index, row in enumerate(item.rows, start=1):
                values: list[str] = []
                for column_index, cell in enumerate(row.cells, start=1):
                    paragraphs: list[str] = []
                    for paragraph_index, paragraph in enumerate(cell.paragraphs):
                        pieces: list[str] = []
                        for inline_index, (event_kind, value) in enumerate(
                            _docx_paragraph_events(document, paragraph)
                        ):
                            if event_kind == "text":
                                pieces.append(str(value))
                                continue
                            image_content, image_name, media_type = value
                            table_image_names.add(image_name)
                            asset = _describe_office_image(
                                image_content,
                                image_name,
                                media_type,
                                {
                                    "source": "docx_table_image",
                                    "row": row_index,
                                    "column": column_index,
                                    "paragraph": paragraph_index,
                                    "inline_order": inline_index,
                                },
                                description_cache,
                            )
                            table_assets.append(asset)
                            pieces.append(asset.marker)
                        paragraphs.append("".join(pieces))
                    values.append("<br>".join(paragraphs).strip())
                rows.append(values)
            if rows:
                header = " | ".join(rows[0])
                blocks.append(
                    ParsedBlock(
                        kind="table",
                        content="\n".join([header, *(" | ".join(row) for row in rows[1:])]),
                        metadata={"table_header": header},
                        section=" / ".join(entry[1] for entry in headings) or None,
                        assets=tuple(table_assets),
                    )
                )
        else:
            # 一个 Word 段落可能同时包含文字和图片。必须按 run 的原始顺序
            # 聚合成一个块，否则图片会把同一段话拆开，检索命中图片时只能返回
            # “图片前后的一小句”，无法还原原文。
            text_parts: list[str] = []
            paragraph_assets: list[ParsedAsset] = []
            style = getattr(item.style, "name", "")
            events = _docx_paragraph_events(document, item)
            for event_kind, value in events:
                if event_kind == "text":
                    paragraph_text = _clean_text(str(value))
                    if paragraph_text:
                        text_parts.append(paragraph_text)
                    continue
                image_content, image_name, media_type = value  # type: ignore[misc]
                referenced_image_names.add(image_name)
                asset = _describe_office_image(
                    image_content,
                    image_name,
                    media_type,
                    {
                        "source": "docx_embedded_image",
                        "paragraph": len(blocks),
                        "inline_order": len(text_parts),
                    },
                    description_cache,
                )
                paragraph_assets.append(asset)
                text_parts.append(asset.marker)
            paragraph_text = _clean_text("".join(text_parts))
            outline_level: int | None = None
            paragraph_properties = getattr(getattr(item, "_p", None), "pPr", None)
            outline = getattr(paragraph_properties, "outlineLvl", None)
            if outline is not None:
                try:
                    outline_level = int(outline.val)
                except (TypeError, ValueError):
                    outline_level = None
            numbering = getattr(paragraph_properties, "numPr", None)
            is_list_paragraph = numbering is not None or bool(
                re.search(r"\b(list|bullet|numbering|编号|项目符号)\b", style, re.IGNORECASE)
            )
            level = _docx_heading_level(
                style,
                paragraph_text,
                outline_level=outline_level,
                allow_numbered_fallback=not is_list_paragraph,
            )
            if level is not None and paragraph_text:
                headings = [entry for entry in headings if entry[0] < level]
                headings.append((level, paragraph_text))
            section = " / ".join(entry[1] for entry in headings) or None
            if paragraph_text:
                blocks.append(
                    ParsedBlock(
                        kind="text",
                        content=paragraph_text,
                        section=section,
                        assets=tuple(paragraph_assets),
                    )
                )
    # OOXML 图片独立存放在 word/media；提取后调用同一视觉模型生成可检索图片说明。
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            for name in archive.namelist():
                if not name.startswith("word/media/") or name.endswith("/"):
                    continue
                if (
                    Path(name).name in table_image_names
                    or Path(name).name in referenced_image_names
                ):
                    # 正文或表格图片已经按原始位置处理，避免重复 OCR 和重复入库。
                    continue
                image_content = archive.read(name)
                media_type = mimetypes.guess_type(name)[0] or "application/octet-stream"
                image_blocks = _parse_office_image(
                    image_content,
                    Path(name).name,
                    media_type,
                    "",
                    None,
                    "docx",
                    description_cache=description_cache,
                )
                blocks.extend(image_blocks)
    except OcrServiceError as exc:
        raise DocumentParseError(str(exc)) from exc
    except (OSError, zipfile.BadZipFile) as exc:
        raise DocumentParseError("DOCX_IMAGE_PARSE_ERROR: Word 内嵌图片无法提取") from exc
    return blocks


def _parse_pdf(content: bytes) -> list[ParsedBlock]:
    """本地优先解析 PDF，仅把需要视觉识别的页面交给千帆。"""
    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise DocumentParseError("PDF_ENCRYPTED: 暂不支持加密 PDF")
        page_texts = [_clean_text(page.extract_text() or "") for page in reader.pages]
        native_text_by_page = _pdf_native_text_blocks(content, page_texts)
        for page_number, text in enumerate(page_texts, start=1):
            if not text and native_text_by_page.get(page_number):
                page_texts[page_number - 1] = "\n\n".join(
                    block.content for block in native_text_by_page[page_number]
                )
        settings = get_settings()
        missing_pages = [index for index, text in enumerate(page_texts, start=1) if not text]
        layout_mode = getattr(settings, "ocr_layout_analysis_mode", "auto")
        if layout_mode not in {"auto", "all"}:
            layout_mode = "auto"
        layout_pages = (
            list(range(1, len(page_texts) + 1)) if layout_mode == "all" else missing_pages
        )
        if (
            getattr(settings, "ocr_layout_analysis_enabled", False)
            and settings.ocr_model_name.strip()
            and layout_pages
        ):
            layouts = analyze_pdf_pages(
                content,
                settings=settings,
                filename="document.pdf",
                page_numbers=layout_pages,
            )
            if layout_mode == "all":
                blocks: list[ParsedBlock] = []
                headings_state: list[tuple[int, str]] = []
                for page_number, layout in sorted(layouts.items()):
                    blocks.extend(
                        _layout_blocks(
                            layout,
                            page_number=page_number,
                            metadata={"format": "pdf_layout"},
                            page_image=content,
                            headings_state=headings_state,
                        )
                    )
                if blocks:
                    return blocks
            else:
                # auto 模式只替换扫描页；有文字层的页面继续使用本地提取结果。
                blocks = []
                headings_state = []
                for page_number, text in enumerate(page_texts, start=1):
                    if text:
                        page_blocks = native_text_by_page.get(page_number) or (
                            _pdf_plain_text_blocks(text, page_number, headings_state)
                        )
                        page_blocks = _rebase_pdf_page_blocks(page_blocks, headings_state)
                        blocks.extend(page_blocks)
                        _sync_pdf_heading_state(headings_state, page_blocks)
                        continue
                    layout = layouts.get(page_number)
                    if layout is None:
                        continue
                    blocks.extend(
                        _layout_blocks(
                            layout,
                            page_number=page_number,
                            metadata={"format": "pdf_layout"},
                            page_image=content,
                            headings_state=headings_state,
                        )
                    )
                if blocks and all(block.content.strip() for block in blocks):
                    return blocks
        if missing_pages:
            if not settings.ocr_model_name.strip():
                raise DocumentParseError(
                    "OCR_REQUIRED: 请在 .env 中配置 KNOWLEDGE_SERVICE_OCR_MODEL_NAME"
                )
            recognized = recognize_pdf_pages(
                content,
                settings=settings,
                filename="document.pdf",
                page_numbers=missing_pages,
            )
            for page_number in missing_pages:
                if page_number not in recognized:
                    raise DocumentParseError(
                        f"OCR_INVALID_RESPONSE: 千帆 OCR 未返回第 {page_number} 页"
                    )
                page_texts[page_number - 1] = _clean_text(recognized[page_number])
        blocks = []
        headings_state = []
        for index, text in enumerate(page_texts, start=1):
            page_blocks = native_text_by_page.get(index) or _pdf_plain_text_blocks(
                text, index, headings_state
            )
            blocks.extend(page_blocks)
            _sync_pdf_heading_state(headings_state, page_blocks)
        native_image_blocks = _pdf_native_image_blocks(content, page_texts, native_text_by_page)
        if native_image_blocks:
            blocks = native_image_blocks
    except DocumentParseError:
        raise
    except OcrServiceError as exc:
        raise DocumentParseError(str(exc)) from exc
    except Exception as exc:
        raise DocumentParseError("PDF_PARSE_ERROR: PDF 文档无法解析") from exc
    if not any(block.content for block in blocks):
        raise DocumentParseError("EMPTY_DOCUMENT: PDF 页内未提取到可用文字")
    return blocks


def _sync_pdf_heading_state(
    headings: list[tuple[int, str]], blocks: list[ParsedBlock]
) -> None:
    """Carry the last explicit chapter path into the next PDF page."""
    section = next((block.section for block in reversed(blocks) if block.section), None)
    if section:
        headings[:] = [
            (level, title)
            for level, title in enumerate(section.split(" / "), start=1)
            if title.strip()
        ]


def _rebase_pdf_page_blocks(
    blocks: list[ParsedBlock], headings: list[tuple[int, str]]
) -> list[ParsedBlock]:
    """Keep native PDF headings in the current path after OCR-only pages."""
    rebased: list[ParsedBlock] = []
    for block in blocks:
        level = block.metadata.get("heading_level")
        try:
            level = int(level) if level is not None else None
        except (TypeError, ValueError):
            level = None
        if level is not None:
            headings[:] = [entry for entry in headings if entry[0] < level]
            headings.append((level, block.content))
        section = _section_path(headings) or block.section
        rebased.append(replace(block, section=section))
    return rebased


def _pdf_plain_text_blocks(
    text: str,
    page_number: int,
    headings: list[tuple[int, str]],
) -> list[ParsedBlock]:
    """Split OCR/PDF fallback text into paragraphs and apply safe numbered headings."""
    blocks: list[ParsedBlock] = []
    for part in re.split(r"\n\s*\n", _clean_text(text)):
        candidate = _clean_text(part)
        if not candidate:
            continue
        inferred = (
            _safe_inferred_heading_level(candidate) if "\n" not in candidate else None
        )
        if inferred is not None:
            headings[:] = [entry for entry in headings if entry[0] < inferred]
            headings.append((inferred, candidate))
        blocks.append(
            ParsedBlock(
                kind="text",
                content=candidate,
                page_number=page_number,
                section=_section_path(headings),
            )
        )
    return blocks


def _pdf_native_text_blocks(
    content: bytes, page_texts: list[str]
) -> dict[int, list[ParsedBlock]]:
    """Use PDF span typography and coordinates to infer conservative heading levels."""
    try:
        import fitz

        document = fitz.open(stream=content, filetype="pdf")
    except Exception:
        return {}
    try:
        page_records: dict[int, list[dict[str, Any]]] = {}
        all_sizes: list[float] = []
        for page_number, page in enumerate(document, start=1):
            records = []
            for reading_order, block in enumerate(page.get_text("dict").get("blocks", [])):
                if block.get("type") != 0:
                    continue
                lines = block.get("lines", [])
                line_texts = [
                    "".join(str(span.get("text", "")) for span in line.get("spans", []))
                    for line in lines
                ]
                text = _clean_text("\n".join(line_texts))
                if not text:
                    continue
                spans = [span for line in lines for span in line.get("spans", [])]
                sizes = [float(span.get("size", 0)) for span in spans if span.get("size")]
                if not sizes:
                    continue
                all_sizes.extend(sizes)
                records.append(
                    {
                        "text": text,
                        "line_count": len([line for line in line_texts if line.strip()]),
                        "size": sum(sizes) / len(sizes),
                        "bold": any(
                            int(span.get("flags", 0)) & 16
                            or "bold" in str(span.get("font", "")).lower()
                            for span in spans
                        ),
                        "bbox": list(block.get("bbox") or (0, 0, 0, 0)),
                        "page_height": float(page.rect.height),
                        "reading_order": reading_order,
                    }
                )
            records.sort(key=lambda item: (item["bbox"][1], item["bbox"][0]))
            page_records[page_number] = records

        if not all_sizes:
            return {}
        ordered_sizes = sorted(all_sizes)
        middle = len(ordered_sizes) // 2
        median_size = (
            ordered_sizes[middle]
            if len(ordered_sizes) % 2
            else (ordered_sizes[middle - 1] + ordered_sizes[middle]) / 2
        )
        candidates: list[dict[str, Any]] = []
        for records in page_records.values():
            for record_index, record in enumerate(records):
                text = record["text"]
                if (
                    record["line_count"] > 2
                    or len(text) > 120
                    or re.search(r"[。！？!?；;:.]\s*$", text)
                ):
                    continue
                numbered_level = _heading_level(text) if "\n" not in text else None
                numbered_subsection = bool(
                    re.match(r"^\s*\d{1,3}[.．]\d", text)
                    or re.match(r"^\s*[（(][一二三四五六七八九十百千万]+[）)]", text)
                )
                prominent = record["size"] >= median_size * 1.2 or (
                    record["bold"] and record["size"] >= median_size * 1.08
                )
                before_gap = (
                    record["bbox"][1] - records[record_index - 1]["bbox"][3]
                    if record_index > 0
                    else record["bbox"][1]
                )
                after_gap = (
                    records[record_index + 1]["bbox"][1] - record["bbox"][3]
                    if record_index + 1 < len(records)
                    else record["page_height"] - record["bbox"][3]
                )
                isolated_position = (
                    before_gap >= median_size * 0.8 and after_gap >= median_size * 0.8
                )
                if (
                    prominent
                    or (numbered_subsection and numbered_level is not None)
                    or (isolated_position and record["size"] >= median_size * 1.05)
                ):
                    candidates.append(record)

        style_sizes = sorted({round(item["size"], 1) for item in candidates}, reverse=True)
        result: dict[int, list[ParsedBlock]] = {}
        headings: list[tuple[int, str]] = []
        for page_number, records in page_records.items():
            output: list[ParsedBlock] = []
            for record in records:
                text = record["text"]
                is_heading = record in candidates
                if is_heading:
                    numbered_level = _heading_level(text) if "\n" not in text else None
                    if numbered_level is not None:
                        level = numbered_level
                    else:
                        size_rank = style_sizes.index(round(record["size"], 1))
                        level = min(6, size_rank + 1)
                    headings[:] = [entry for entry in headings if entry[0] < level]
                    headings.append((level, text))
                output.append(
                    ParsedBlock(
                        kind="text",
                        content=text,
                        page_number=page_number,
                        section=_section_path(headings),
                        metadata={
                            "format": "pdf_native",
                            "reading_order": record["reading_order"],
                            "bbox": record["bbox"],
                            "heading_level": headings[-1][0]
                            if is_heading and headings
                            else None,
                        },
                    )
                )
            if not output and page_number <= len(page_texts) and page_texts[page_number - 1]:
                output = _pdf_plain_text_blocks(page_texts[page_number - 1], page_number, headings)
            result[page_number] = output
        return result
    except Exception:
        return {}
    finally:
        document.close()


def _pdf_native_image_blocks(
    content: bytes,
    page_texts: list[str],
    native_text_by_page: dict[int, list[ParsedBlock]] | None = None,
) -> list[ParsedBlock]:
    """按 PDF 原生文字块和图片坐标重建阅读顺序，避免有文字层的图片被漏掉。"""
    try:
        import fitz

        document = fitz.open(stream=content, filetype="pdf")
    except Exception:
        return []
    try:
        page_data: list[tuple[int, Any, list[dict[str, Any]]]] = []
        native_text_by_page = native_text_by_page or _pdf_native_text_blocks(content, page_texts)
        images: list[tuple[bytes, str, str]] = []
        for page_number, page in enumerate(document, start=1):
            raw_blocks = page.get_text("dict").get("blocks", [])
            if not any(block.get("type") == 1 and block.get("image") for block in raw_blocks):
                continue
            page_data.append((page_number, page, raw_blocks))
            for block in raw_blocks:
                if block.get("type") != 1 or not block.get("image"):
                    continue
                name = f"page-{page_number}.{block.get('ext') or 'png'}"
                media_type = mimetypes.guess_type(name)[0] or "image/png"
                images.append((bytes(block["image"]), name, media_type))
        if not page_data:
            return []
        descriptions: dict[str, str] = {}
        _prewarm_office_image_descriptions(images, descriptions, "pdf")
        by_page: dict[int, list[ParsedBlock]] = {}
        for page_number, page, raw_blocks in page_data:
            ordered = sorted(
                raw_blocks,
                key=lambda block: (
                    float((block.get("bbox") or (0, 0, 0, 0))[1]),
                    float((block.get("bbox") or (0, 0, 0, 0))[0]),
                ),
            )
            output: list[ParsedBlock] = []
            native_text_found = False
            page_text_blocks = native_text_by_page.get(page_number, [])
            section_by_bbox = {
                tuple(round(float(value), 1) for value in block.metadata.get("bbox", [])): block
                for block in page_text_blocks
                if block.metadata.get("bbox")
            }
            current_section = next(
                (
                    block.section
                    for prior_page in range(page_number - 1, 0, -1)
                    for block in reversed(native_text_by_page.get(prior_page, []))
                    if block.section
                ),
                None,
            )
            for reading_order, block in enumerate(ordered):
                bbox = block.get("bbox") or (0, 0, 0, 0)
                if block.get("type") == 0:
                    lines = [
                        "".join(span.get("text", "") for span in line.get("spans", []))
                        for line in block.get("lines", [])
                    ]
                    text = _clean_text("\n".join(lines))
                    if text:
                        native_text_found = True
                        source_block = section_by_bbox.get(
                            tuple(round(float(value), 1) for value in bbox)
                        )
                        if source_block and source_block.section:
                            current_section = source_block.section
                        output.append(
                            ParsedBlock(
                                kind="text",
                                content=text,
                                page_number=page_number,
                                section=(source_block.section if source_block else None)
                                or current_section,
                                metadata={
                                    "format": "pdf_native",
                                    "reading_order": reading_order,
                                    "bbox": list(bbox),
                                    **(
                                        {"heading_level": source_block.metadata["heading_level"]}
                                        if source_block
                                        and source_block.metadata.get("heading_level") is not None
                                        else {}
                                    ),
                                },
                            )
                        )
                elif block.get("type") == 1 and block.get("image"):
                    image_bytes = bytes(block["image"])
                    extension = str(block.get("ext") or "png")
                    name = f"page-{page_number}-{reading_order}.{extension}"
                    media_type = mimetypes.guess_type(name)[0] or "image/png"
                    normalized_bbox = [
                        float(bbox[0]) / page.rect.width * 1000,
                        float(bbox[1]) / page.rect.height * 1000,
                        float(bbox[2]) / page.rect.width * 1000,
                        float(bbox[3]) / page.rect.height * 1000,
                    ]
                    marker = _asset_marker(
                        image_bytes,
                        name,
                        "pdf_native",
                        page=page_number,
                        order=reading_order,
                        bbox=normalized_bbox,
                    )
                    description = descriptions.get(
                        hashlib.sha256(image_bytes).hexdigest(), ""
                    )
                    output.append(
                        ParsedBlock(
                            kind="text",
                            content=f"{description}\n{marker}" if description else marker,
                            page_number=page_number,
                            section=current_section,
                            metadata={
                                "format": "pdf_native",
                                "reading_order": reading_order,
                                "bbox": normalized_bbox,
                                "asset_marker": marker,
                                "asset_description": description,
                            },
                            asset_content=image_bytes,
                            asset_media_type=media_type,
                        )
                    )
            if (
                not native_text_found
                and page_number <= len(page_texts)
                and page_texts[page_number - 1]
            ):
                output[0:0] = page_text_blocks or _pdf_plain_text_blocks(
                    page_texts[page_number - 1], page_number, []
                )
            by_page[page_number] = output
        result = []
        for page_number, text in enumerate(page_texts, start=1):
            result.extend(
                by_page.get(page_number)
                or native_text_by_page.get(page_number)
                or [ParsedBlock(kind="text", content=text, page_number=page_number)]
            )
        return result
    except Exception:
        return []
    finally:
        document.close()
