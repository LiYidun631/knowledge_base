"""验证解析和父子切分保持正文、结构与稳定标识。"""

import io
import threading
import time
import zipfile
from types import SimpleNamespace
from uuid import uuid4

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pypdf import PdfWriter

from knowledge_service.application.documents import parsing
from knowledge_service.application.documents.chunking import build_chunks
from knowledge_service.application.documents.parsing import (
    DocumentParseError,
    ParsedBlock,
    parse_document,
)
from knowledge_service.core.config import Settings
from knowledge_service.infrastructure.ocr_client import LayoutBlock, LayoutResult


def test_child_overlap_preserves_all_source_text() -> None:
    """超长文本跨子块时不能因添加重叠而丢掉任何字符。"""
    text = "".join(chr(0x4E00 + index) for index in range(1600))
    document_id = uuid4()
    chunks = build_chunks(document_id, [ParsedBlock(kind="text", content=text, page_number=3)])
    children = [chunk for chunk in chunks if chunk.metadata["level"] == "child"]
    assert all(len(chunk.content) <= 600 for chunk in children)
    assert all(chunk.page_number == 3 for chunk in children)
    assert all(character in "".join(chunk.content for chunk in children) for character in text)
    assert chunks == build_chunks(
        document_id, [ParsedBlock(kind="text", content=text, page_number=3)]
    )


def test_csv_split_keeps_header_and_every_row() -> None:
    """所有表格子块应带列头，完整保留每个数据行。"""
    rows = [f"patient-{index:04d},诊断内容{index}" for index in range(100)]
    blocks = parse_document("data.csv", ("姓名,诊断\n" + "\n".join(rows)).encode())
    children = [
        chunk for chunk in build_chunks(uuid4(), blocks) if chunk.metadata["level"] == "child"
    ]
    assert len(children) > 1
    assert all(chunk.content.startswith("姓名 | 诊断\n") for chunk in children)
    assert all(len(chunk.content) <= 600 for chunk in children)
    joined = "\n".join(chunk.content for chunk in children)
    assert all(f"patient-{index:04d}" in joined for index in range(100))


def test_docx_preserves_paragraph_table_order_and_section() -> None:
    """Word 表格前后的段落应保持原始顺序与章节来源。"""
    word = Document()
    word.add_heading("诊疗方案", level=1)
    word.add_paragraph("表格前的说明")
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "结果"
    table.cell(1, 0).text = "血压"
    table.cell(1, 1).text = "正常"
    word.add_paragraph("表格后的说明")
    stream = io.BytesIO()
    word.save(stream)
    blocks = parse_document("说明.docx", stream.getvalue())
    assert [block.kind for block in blocks] == ["text", "text", "table", "text"]
    assert blocks[-1].section == "诊疗方案"
    assert "项目 | 结果" in blocks[2].content


def test_docx_numbered_normal_paragraphs_create_nested_sections(monkeypatch) -> None:
    """说明书把编号小节存成 Normal 时，后续段落仍应归入正确章节。"""
    word = Document()
    word.add_heading("3 正确的使用方法", level=1)
    word.add_paragraph("3.1 测量步骤")
    word.add_paragraph("步骤1：装好电池并开机。")
    word.add_paragraph("3.1.1 耳温的测量方法")
    word.add_paragraph("将探头轻轻插入耳道。")
    word.add_paragraph("3.1.2 额温的测量方法")
    word.add_paragraph("将探头对准额头中心。")
    word.add_paragraph("3.2 查看记忆")
    word.add_paragraph("短按电源键查看上次测量结果。")
    stream = io.BytesIO()
    word.save(stream)
    monkeypatch.setattr(
        parsing,
        "get_settings",
        lambda: Settings(_env_file=None, ocr_embedded_images_enabled=False),
    )

    blocks = parse_document("guide.docx", stream.getvalue())

    assert blocks[2].section == "3 正确的使用方法 / 3.1 测量步骤"
    assert blocks[4].section == ("3 正确的使用方法 / 3.1 测量步骤 / 3.1.1 耳温的测量方法")
    assert blocks[6].section == ("3 正确的使用方法 / 3.1 测量步骤 / 3.1.2 额温的测量方法")
    assert blocks[8].section == "3 正确的使用方法 / 3.2 查看记忆"


@pytest.mark.parametrize(
    ("title", "expected_level"),
    [
        ("① 测量准备", 1),
        ("一、测量准备", 1),
        ("（一）测量准备", 2),
        ("(1) 测量准备", 2),
        ("2.1 测量步骤", 2),
    ],
)
def test_docx_heading_number_formats_are_supported(title, expected_level) -> None:
    assert parsing._docx_heading_level("Normal", title) == expected_level


@pytest.mark.parametrize(
    "text",
    [
        "①按下电源键。",
        "（1）等待提示音。",
        "一、将探头插入耳道后开始测量。",
        "步骤1：装好电池并开机。",
    ],
)
def test_docx_numbered_steps_are_not_headings(text) -> None:
    assert parsing._docx_heading_level("Normal", text) is None


def test_docx_native_list_paragraphs_skip_numbered_heading_fallback() -> None:
    assert (
        parsing._docx_heading_level(
            "List Number", "1. Press the power button", allow_numbered_fallback=False
        )
        is None
    )


def test_markdown_recognizes_setext_and_ignores_code_fence_headings() -> None:
    source = (
        "Guide\n=====\n\nIntroduction text.\n\n"
        "## Details\n\nDetails text.\n\n"
        "```markdown\n# literal code heading\n```\n\nMore details."
    )

    blocks = parse_document("guide.md", source.encode())

    assert [block.section for block in blocks] == [
        "Guide",
        "Guide / Details",
        "Guide / Details",
        "Guide / Details",
    ]
    assert blocks[2].content == "```markdown\n# literal code heading\n```"
    assert blocks[1].metadata["section_path"] == ["Guide", "Details"]
    assert blocks[1].metadata["section_path_ids"] == blocks[3].metadata["section_path_ids"]


def test_markdown_numbered_steps_are_not_inferred_as_sections() -> None:
    blocks = parse_document(
        "steps.md",
        b"1. Press the power button.\n\n2. Wait for the result.",
    )

    assert [block.section for block in blocks] == [None, None]


def test_ocr_heading_blocks_form_nested_sections() -> None:
    layout = LayoutResult(
        (
            LayoutBlock("heading", "Guide", metadata={"heading_level": 1}),
            LayoutBlock("heading", "Setup", metadata={"heading_level": 2}),
            LayoutBlock("text", "Press the button."),
        ),
        "Guide\n\nSetup\n\nPress the button.",
    )

    blocks = parsing._layout_blocks(layout, page_number=1)

    assert [block.section for block in blocks] == [
        "Guide",
        "Guide / Setup",
        "Guide / Setup",
    ]


def test_docx_table_image_stays_in_its_cell(monkeypatch) -> None:
    """DOCX 表格内图片应绑定到对应表格块，而不是追加到文档末尾。"""
    image_bytes = (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01"
        b"\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
        b"\x00\x00\x00\x0dIDAT\x08\xd7c\xf8\xcf\xc0\xf0\x1f\x00\x05\x00\x01\xff\x89\x99='\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    word = Document()
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "检查项目"
    table.cell(0, 1).text = "结果"
    table.cell(1, 0).text = "化验单"
    table.cell(1, 1).paragraphs[0].add_run().add_picture(io.BytesIO(image_bytes))
    stream = io.BytesIO()
    word.save(stream)
    monkeypatch.setattr(
        parsing,
        "analyze_image",
        lambda content, *, settings, filename: LayoutResult(
            (LayoutBlock("image", "化验结果图片"),), "化验结果图片"
        ),
    )
    monkeypatch.setattr(
        parsing,
        "get_settings",
        lambda: Settings(
            _env_file=None,
            ocr_embedded_image_min_bytes=0,
            ocr_embedded_image_min_width=1,
            ocr_embedded_image_min_height=1,
        ),
    )

    blocks = parse_document("table.docx", stream.getvalue())

    assert len(blocks) == 1
    assert blocks[0].kind == "table"
    assert "[[KB_ASSET:" in blocks[0].content
    assert len(blocks[0].assets) == 1
    assert blocks[0].assets[0].metadata["source"] == "docx_table_image"
    assert blocks[0].assets[0].metadata["row"] == 2
    assert blocks[0].assets[0].metadata["column"] == 2
    assert "asset_marker" in blocks[0].assets[0].metadata


def test_docx_images_use_bounded_concurrency_and_isolate_failures(monkeypatch) -> None:
    """DOCX 图片并发识别受配置限流，单张失败不影响其他图片和正文顺序。"""

    def make_image(color: str) -> bytes:
        stream = io.BytesIO()
        Image.new("RGB", (100, 80), color).save(stream, format="PNG")
        return stream.getvalue()

    failed_image = make_image("red")
    successful_image = make_image("blue")
    word = Document()
    word.add_heading("工作原理", level=1)
    paragraph = word.add_paragraph()
    paragraph.add_run("第一张图前的说明")
    paragraph.add_run().add_picture(io.BytesIO(failed_image))
    paragraph.add_run("第一张图后的说明")
    word.add_paragraph().add_run().add_picture(io.BytesIO(successful_image))
    stream = io.BytesIO()
    word.save(stream)

    lock = threading.Lock()
    active_requests = 0
    max_active_requests = 0

    def fake_analyze(image, *, settings, filename):
        nonlocal active_requests, max_active_requests
        with lock:
            active_requests += 1
            max_active_requests = max(max_active_requests, active_requests)
        try:
            time.sleep(0.05)
            if image == failed_image:
                raise parsing.OcrServiceError("OCR_TIMEOUT: 模拟单张图片超时")
            return LayoutResult((LayoutBlock("text", "第二张图片识别结果"),), "第二张图片识别结果")
        finally:
            with lock:
                active_requests -= 1

    monkeypatch.setattr(parsing, "analyze_image", fake_analyze)
    monkeypatch.setattr(
        parsing,
        "get_settings",
        lambda: Settings(
            _env_file=None,
            ocr_embedded_concurrency=2,
            ocr_embedded_image_min_bytes=0,
            ocr_embedded_image_min_width=1,
            ocr_embedded_image_min_height=1,
        ),
    )

    blocks = parse_document("guide.docx", stream.getvalue())
    contents = [block.content for block in blocks]

    assert max_active_requests == 2
    assert any("[[KB_ASSET:" in content for content in contents)
    # 同一 Word 段落中的文字和图片保留为一个块，图片占位符仍处于原文位置。
    paragraph_content = next(content for content in contents if "[[KB_ASSET:" in content)
    assert paragraph_content.index("第一张图前的说明") < paragraph_content.index("[[KB_ASSET:")
    assert paragraph_content.index("第一张图后的说明") > paragraph_content.index("[[KB_ASSET:")
    assert sum(len(block.assets) + int(block.asset_content is not None) for block in blocks) == 2


def test_pdf_without_text_requires_ocr(monkeypatch) -> None:
    """无文字的 PDF 不应当被误判为已经完成入库。"""
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    stream = io.BytesIO()
    writer.write(stream)
    monkeypatch.setattr(
        parsing,
        "get_settings",
        lambda: SimpleNamespace(ocr_model_name=""),
    )
    with pytest.raises(DocumentParseError, match="OCR_REQUIRED"):
        parse_document("scan.pdf", stream.getvalue())


def test_scanned_pdf_uses_qianfan_and_keeps_page_number(monkeypatch) -> None:
    """扫描 PDF 只将无文字页交给千帆识别，并保留原始页码。"""
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    stream = io.BytesIO()
    writer.write(stream)
    pdf_bytes = stream.getvalue()
    settings = SimpleNamespace(ocr_model_name="qianfan-ocr")
    calls = []

    def fake_recognize(content, *, settings, filename, page_numbers):
        calls.append((content, settings, filename, page_numbers))
        return {1: "扫描报告正文"}

    monkeypatch.setattr(parsing, "get_settings", lambda: settings)
    monkeypatch.setattr(parsing, "recognize_pdf_pages", fake_recognize)

    blocks = parse_document("scan.pdf", pdf_bytes)

    assert [(block.content, block.page_number) for block in blocks] == [("扫描报告正文", 1)]
    assert calls == [(pdf_bytes, settings, "document.pdf", [1])]


def test_text_pdf_does_not_call_vision_model_in_auto_mode(monkeypatch) -> None:
    """有文字层的 PDF 在 auto 模式下直接本地解析，避免逐页调用视觉模型。"""
    settings = SimpleNamespace(
        ocr_model_name="qianfan-ocr",
        ocr_layout_analysis_enabled=True,
        ocr_layout_analysis_mode="auto",
    )
    page = SimpleNamespace(extract_text=lambda: "本地文字层")
    reader = SimpleNamespace(is_encrypted=False, pages=[page])
    monkeypatch.setattr(parsing, "get_settings", lambda: settings)
    monkeypatch.setattr(parsing, "PdfReader", lambda *_args, **_kwargs: reader)
    monkeypatch.setattr(
        parsing,
        "analyze_pdf_pages",
        lambda *args, **kwargs: pytest.fail("文字页不应调用视觉模型"),
    )

    blocks = parse_document("text.pdf", b"pdf-bytes")

    assert [(block.content, block.page_number) for block in blocks] == [("本地文字层", 1)]


def test_native_pdf_uses_typography_for_heading_paths_across_pages(monkeypatch) -> None:
    import fitz

    document = fitz.open()
    first = document.new_page()
    first.insert_text((72, 60), "Operations Guide", fontsize=22, fontname="hebo")
    first.insert_text((72, 100), "3.1 Operation", fontsize=16, fontname="hebo")
    first.insert_text((72, 140), "Press the power button.", fontsize=11)
    first.insert_text((72, 160), "Wait for the result.", fontsize=11)
    second = document.new_page()
    second.insert_text((72, 72), "Hold the power button to stop.", fontsize=11)
    pdf_bytes = document.tobytes()
    document.close()
    monkeypatch.setattr(
        parsing,
        "get_settings",
        lambda: SimpleNamespace(
            ocr_model_name="",
            ocr_layout_analysis_enabled=False,
            ocr_layout_analysis_mode="auto",
        ),
    )

    blocks = parse_document("guide.pdf", pdf_bytes)

    body = next(block for block in blocks if block.content == "Press the power button.")
    later_page = next(
        block for block in blocks if block.content == "Hold the power button to stop."
    )
    assert body.section == "Operations Guide / 3.1 Operation"
    assert later_page.section == body.section
    assert body.metadata["heading_level"] is None


def test_native_pdf_recovers_vector_table_and_drops_duplicate_text(monkeypatch) -> None:
    """带文字层的矢量线框表格应还原为表格块，且单元格文字不再重复成正文。"""
    import fitz

    document = fitz.open()
    page = document.new_page(width=400, height=300)
    xs = [60, 160, 260, 360]
    ys = [80, 130, 180]
    for column in xs:
        page.draw_line((column, ys[0]), (column, ys[-1]))
    for row in ys:
        page.draw_line((xs[0], row), (xs[-1], row))
    cells = [["型号", "重量", "单位"], ["HBF-214", "1.6", "kg"]]
    for row_index, row in enumerate(cells):
        for column_index, text in enumerate(row):
            page.insert_text((xs[column_index] + 6, ys[row_index] + 30), text, fontsize=12)
    pdf_bytes = document.tobytes()
    document.close()
    monkeypatch.setattr(
        parsing,
        "get_settings",
        lambda: SimpleNamespace(
            ocr_model_name="",
            ocr_layout_analysis_enabled=False,
            ocr_layout_analysis_mode="auto",
        ),
    )

    blocks = parse_document("spec.pdf", pdf_bytes)

    table_blocks = [block for block in blocks if block.kind == "table"]
    assert len(table_blocks) == 1
    table = table_blocks[0]
    assert table.metadata["table_header"]
    assert "HBF-214" in table.content
    # 表格内的文字不应再作为普通正文块重复出现。
    assert all("HBF-214" not in block.content for block in blocks if block.kind == "text")


def test_text_pdf_crops_vision_located_vector_illustrations(monkeypatch) -> None:
    """有文字层且含矢量绘图的页面只向视觉模型要图片区域，大图按位置插入正文。"""
    import fitz

    document = fitz.open()
    page = document.new_page(width=400, height=300)
    page.insert_text((60, 40), "Step one", fontsize=12)
    page.draw_rect(fitz.Rect(100, 100, 300, 260))
    page.insert_text((60, 280), "Step two", fontsize=12)
    pdf_bytes = document.tobytes()
    document.close()

    layout = LayoutResult(
        blocks=(
            LayoutBlock(kind="image", text="图片", bbox=(250.0, 333.0, 750.0, 866.0)),
            # 与上图重叠超过七成，属同一示意图的重复检测，应被去重丢弃。
            LayoutBlock(kind="image", text="图片", bbox=(260.0, 340.0, 740.0, 860.0)),
            # 按钮小图标达不到示意图尺寸，应被过滤掉。
            LayoutBlock(kind="image", text="图片", bbox=(10.0, 10.0, 60.0, 60.0)),
        ),
        markdown="",
    )
    calls = []

    def fake_analyze_pdf_pages(content, *, settings, filename, page_numbers):
        calls.append(page_numbers)
        return {page_number: layout for page_number in page_numbers}

    monkeypatch.setattr(parsing, "analyze_pdf_pages", fake_analyze_pdf_pages)

    def fake_describe(content, name, media_type, metadata, description_cache=None):
        return parsing.ParsedAsset(content, media_type, name, "示意图", metadata)

    monkeypatch.setattr(parsing, "_describe_office_image", fake_describe)
    monkeypatch.setattr(
        parsing,
        "get_settings",
        lambda: SimpleNamespace(
            ocr_model_name="vision",
            ocr_layout_analysis_enabled=True,
            ocr_layout_analysis_mode="auto",
            ocr_embedded_concurrency=3,
        ),
    )

    blocks = parse_document("foldout.pdf", pdf_bytes)

    assert calls == [[1]]
    image_blocks = [block for block in blocks if block.asset_content]
    assert len(image_blocks) == 1
    image = image_blocks[0]
    assert image.page_number == 1
    assert image.content.startswith("[[KB_ASSET:")
    assert image.metadata["asset_description"] == "示意图"
    assert image.metadata["asset_marker"] == image.content
    assert [round(value, 1) for value in image.metadata["bbox"]] == [
        250.0,
        333.0,
        750.0,
        866.0,
    ]
    texts = [block for block in blocks if block.kind == "text" and not block.asset_content]
    assert texts[0].content == "Step one"
    assert texts[-1].content == "Step two"


def test_pure_text_page_skips_vision_image_detection(monkeypatch) -> None:
    """纯文字页面没有矢量内容，不应触发视觉模型调用。"""
    import fitz

    document = fitz.open()
    page = document.new_page(width=400, height=300)
    page.insert_text((60, 40), "Plain paragraph.", fontsize=12)
    pdf_bytes = document.tobytes()
    document.close()
    monkeypatch.setattr(
        parsing,
        "get_settings",
        lambda: SimpleNamespace(
            ocr_model_name="vision",
            ocr_layout_analysis_enabled=True,
            ocr_layout_analysis_mode="auto",
        ),
    )
    monkeypatch.setattr(
        parsing,
        "analyze_pdf_pages",
        lambda *args, **kwargs: pytest.fail("纯文字页不应调用视觉模型"),
    )

    blocks = parse_document("text.pdf", pdf_bytes)

    assert [block.content for block in blocks] == ["Plain paragraph."]


def test_native_pdf_reads_two_column_layout_in_reading_order(monkeypatch) -> None:
    """双栏版面应按「左栏自上而下，再右栏」输出；同基线被合并成宽块的也要拆回去。"""
    import fitz

    document = fitz.open()
    page = document.new_page(width=500, height=400)
    # 故意打乱写入顺序，模拟 PDF 内容流与视觉顺序不一致。
    page.insert_text((320, 120), "Right column top", fontsize=12)
    page.insert_text((60, 120), "Left column top", fontsize=12)
    page.insert_text((320, 220), "Right column bottom", fontsize=12)
    page.insert_text((60, 220), "Left column bottom", fontsize=12)
    pdf_bytes = document.tobytes()
    document.close()
    monkeypatch.setattr(
        parsing,
        "get_settings",
        lambda: SimpleNamespace(
            ocr_model_name="",
            ocr_layout_analysis_enabled=False,
            ocr_layout_analysis_mode="auto",
        ),
    )

    blocks = parse_document("foldout.pdf", pdf_bytes)

    texts = [block.content for block in blocks if block.kind == "text"]
    assert texts == [
        "Left column top",
        "Left column bottom",
        "Right column top",
        "Right column bottom",
    ]


def test_html_removes_scripts_and_avoids_duplicate_table_cells() -> None:
    """HTML 的脚本和表格内部段落不能重复进入正文。"""
    source = (
        "<h1>指南</h1><script>secret</script><p>正文</p>"
        "<table><tr><th>列名</th></tr><tr><td><p>数据</p></td></tr></table>"
    )
    blocks = parse_document("guide.html", source.encode())
    assert [block.kind for block in blocks] == ["text", "table"]
    assert all("secret" not in block.content for block in blocks)
    assert blocks[0].section == "指南"


def test_json_is_pretty_printed_for_retrieval() -> None:
    """JSON 应保留中文和层级，便于后续分块与检索。"""
    blocks = parse_document("settings.json", '{"name":"知识库","enabled":true}'.encode())
    assert blocks[0].metadata["format"] == "json"
    assert '"name": "知识库"' in blocks[0].content
    assert '"enabled": true' in blocks[0].content


def test_xlsx_preserves_sheet_name_and_rows() -> None:
    """XLSX 应保留工作表名称和表格行。"""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "患者资料"
    worksheet.append(["姓名", "诊断"])
    worksheet.append(["张三", "血压正常"])
    stream = io.BytesIO()
    workbook.save(stream)

    blocks = parse_document("patients.xlsx", stream.getvalue())

    assert blocks[0].section == "患者资料"
    assert "姓名 | 诊断" in blocks[0].content
    assert "张三 | 血压正常" in blocks[0].content


def test_pptx_preserves_slide_text() -> None:
    """PPTX 应把幻灯片文字转换为带页码的文本块。"""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(0, 0, 5000000, 1000000)
    textbox.text = "健康管理指南"
    stream = io.BytesIO()
    presentation.save(stream)

    blocks = parse_document("guide.pptx", stream.getvalue())

    assert blocks[0].page_number == 1
    assert blocks[0].section == "第 1 页"
    assert blocks[0].content == "健康管理指南"


def test_pptx_extracts_embedded_image_for_ocr(monkeypatch) -> None:
    """PPTX 内嵌图片应生成可检索说明并保留原图资产。"""
    image_bytes = (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01"
        b"\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
        b"\x00\x00\x00\x0dIDAT\x08\xd7c\xf8\xcf\xc0\xf0\x1f\x00\x05\x00\x01\xff\x89\x99='\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(image_bytes), 0, 0)
    stream = io.BytesIO()
    presentation.save(stream)
    monkeypatch.setattr(
        parsing,
        "analyze_image",
        lambda content, *, settings, filename: LayoutResult(
            (LayoutBlock("image", "医学示意图"),), "医学示意图"
        ),
    )
    monkeypatch.setattr(
        parsing,
        "get_settings",
        lambda: Settings(
            _env_file=None,
            ocr_embedded_image_min_bytes=0,
            ocr_embedded_image_min_width=1,
            ocr_embedded_image_min_height=1,
        ),
    )

    blocks = parse_document("guide.pptx", stream.getvalue())

    assert "[[KB_ASSET:" in blocks[0].content
    assert blocks[0].asset_content == image_bytes


def test_legacy_doc_requires_office_converter(monkeypatch) -> None:
    """旧版 DOC 没有转换器时应返回可操作的错误，而不是静默失败。"""
    monkeypatch.setattr(parsing, "_find_office_converter", lambda: None)

    with pytest.raises(DocumentParseError, match="OFFICE_CONVERTER_NOT_CONFIGURED"):
        parse_document("legacy.doc", b"not-a-real-doc")


def test_default_extensions_include_common_formats() -> None:
    """默认配置应覆盖常见文本、办公文件和图片格式。"""
    settings = Settings(_env_file=None)
    assert {".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx", ".jpg"}.issubset(
        settings.allowed_document_extensions
    )


def test_invalid_chunk_parameters_rejected() -> None:
    """阻止无效的重叠与分块长度配置。"""
    with pytest.raises(ValueError, match="子块重叠"):
        Settings(child_chunk_chars=100, child_chunk_overlap_chars=100)


def _markdown_table_source() -> bytes:
    return (
        "# 检验报告\n\n"
        "| 项目 | 结果 | 参考范围 |\n"
        "| --- | --- | --- |\n"
        "| 白细胞 | 6.5 | 4.0-10.0 |\n"
        "| 血小板 | 210 | 100-300 |\n\n"
        "表格后的正文。\n"
    ).encode()


def test_markdown_table_becomes_table_block_with_header() -> None:
    """GFM 表格应识别为表格块，子块重复列头且不丢数据行。"""
    blocks = parse_document("report.md", _markdown_table_source())
    table = next(block for block in blocks if block.kind == "table")
    assert table.metadata["table_header"] == "项目 | 结果 | 参考范围"
    assert "白细胞 | 6.5 | 4.0-10.0" in table.content
    assert table.section == "检验报告"
    children = [
        chunk
        for chunk in build_chunks(uuid4(), blocks)
        if chunk.metadata["level"] == "child" and chunk.metadata["block_type"] == "table"
    ]
    assert children and all(
        chunk.content.startswith("项目 | 结果 | 参考范围") for chunk in children
    )


def test_markdown_code_fence_pipes_are_not_tables() -> None:
    """代码围栏里的竖线行不能误判为表格。"""
    source = ("示例代码：\n\n```text\n| a | b |\n| --- | --- |\n|x|y|\n```\n\n结束。\n").encode()
    blocks = parse_document("code.md", source)
    assert all(block.kind == "text" for block in blocks)
    assert "| a | b |" in "".join(block.content for block in blocks)


def test_html_br_keeps_line_breaks() -> None:
    """HTML 段落内的 br 应保留换行，避免两行文字粘连。"""
    source = (
        b"<html><body><p>"
        b"\xe7\xac\xac\xe4\xb8\x80\xe8\xa1\x8c<br>"
        b"\xe7\xac\xac\xe4\xba\x8c\xe8\xa1\x8c</p></body></html>"
    )
    blocks = parse_document("page.html", source)
    assert "第一行\n第二行" in blocks[0].content


def _epub_bytes() -> bytes:
    """构造最小 EPUB2：两章 XHTML 按 spine 顺序解析。"""
    container = (
        '<?xml version="1.0"?>'
        '<container version="1.0" '
        'xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
        '<rootfiles><rootfile full-path="OEBPS/content.opf" '
        'media-type="application/oebps-package+xml"/>'
        "</rootfiles></container>"
    )
    opf = (
        '<?xml version="1.0"?>'
        '<package xmlns="http://www.idpf.org/2007/opf" version="2.0">'
        "<manifest>"
        '<item id="ch1" href="ch1.xhtml" media-type="application/xhtml+xml"/>'
        '<item id="ch2" href="ch2.xhtml" media-type="application/xhtml+xml"/>'
        "</manifest>"
        '<spine><itemref idref="ch2"/><itemref idref="ch1"/></spine>'
        "</package>"
    )
    ch1 = "<html><body><h1>产品概述</h1><p>适用于家用场景。<br>支持蓝牙连接。</p></body></html>"
    ch2 = "<html><body><h2>操作步骤</h2><p>长按电源键开机。</p></body></html>"
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("mimetype", "application/epub+zip")
        archive.writestr("META-INF/container.xml", container)
        archive.writestr("OEBPS/content.opf", opf)
        archive.writestr("OEBPS/ch1.xhtml", ch1)
        archive.writestr("OEBPS/ch2.xhtml", ch2)
    return stream.getvalue()


def test_epub_parses_spine_order_and_sections() -> None:
    """EPUB 应按 spine 顺序解析，章节标题并入章节路径，br 换行保留。"""
    blocks = parse_document("manual.epub", _epub_bytes())
    sections = [block.section for block in blocks]
    assert any(section and section.startswith("操作步骤") for section in sections)
    assert any(section and section.startswith("产品概述") for section in sections)
    assert sections[0] is not None and sections[0].startswith("操作步骤")
    assert "长按电源键开机。" in "".join(block.content for block in blocks)
    assert "支持蓝牙连接。" in "".join(block.content for block in blocks)
